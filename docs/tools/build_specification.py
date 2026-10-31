"""Build the Word specification from the three reviewed Markdown source documents.

Requires python-docx. Rendering and visual review remain a separate verification step.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCES = [
    ROOT / "01-商业化白皮书与商业计划书(Commercial_BP).md",
    ROOT / "02-产品需求规格说明书(PRD_and_UX_Spec).md",
    ROOT / "03-系统技术架构与数据字典规范(Technical_Architecture).md",
]
OUTPUT = ROOT / "Looma(灵犀胶囊)_ToB与ToC商业化产品全景规格书.docx"
NAVY = "183344"
TEAL = "136E6A"
GREY = "60717A"


def east_asia_font(style, name="宋体"):
    style.font.name = name
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def inline(paragraph, text):
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^\)]+\))")
    for part in pattern.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(TEAL)
        elif part.startswith("["):
            match = re.fullmatch(r"\[([^\]]+)\]\((.+)\)", part)
            if match:
                hyperlink = OxmlElement("w:hyperlink")
                rel = paragraph.part.relate_to(
                    match[2],
                    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                    is_external=True,
                )
                hyperlink.set(qn("r:id"), rel)
                run = OxmlElement("w:r")
                props = OxmlElement("w:rPr")
                colour = OxmlElement("w:color")
                colour.set(qn("w:val"), TEAL)
                props.append(colour)
                run.append(props)
                value = OxmlElement("w:t")
                value.text = match[1]
                run.append(value)
                hyperlink.append(run)
                paragraph._p.append(hyperlink)
            else:
                paragraph.add_run(part)
        else:
            paragraph.add_run(part)


def add_table(document, rows):
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    widths = [16.7 / len(rows[0])] * len(rows[0])
    if len(rows[0]) == 3:
        widths = [3.2, 6.7, 6.8]
    elif len(rows[0]) == 4:
        widths = [3.0, 4.1, 5.0, 4.6]
    for index, values in enumerate(rows):
        cells = table.rows[0].cells if index == 0 else table.add_row().cells
        row = table.rows[index]
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        if index == 0:
            repeat = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(repeat)
        for col, value in enumerate(values):
            cell = cells[col]
            cell.width = Cm(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.15
            # Keep a repeated header with the first body row, not alone at a page foot.
            paragraph.paragraph_format.keep_with_next = index == 0
            inline(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(9)
                if index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(NAVY)
            if index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E4F0ED")
                cell._tc.get_or_add_tcPr().append(shade)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_markdown(document, source):
    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            index += 1
            continue
        if line.startswith("```"):
            language = line[3:]
            index += 1
            block = []
            while index < len(lines) and not lines[index].startswith("```"):
                block.append(lines[index])
                index += 1
            if language == "mermaid":
                paragraph = document.add_paragraph(style="Flow")
                paragraph.add_run("输入与本地草稿 → 四类核对 → 原子收尾 → 项目接续与提醒")
                paragraph = document.add_paragraph(style="Flow")
                paragraph.add_run("可选模型仅提候选；收尾保留来源与撤销；选定信息另建纪要审核。")
            else:
                for code_line in block:
                    paragraph = document.add_paragraph(style="Code")
                    paragraph.add_run(code_line)
            index += 1
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                    rows.append(cells)
                index += 1
            if rows:
                add_table(document, rows)
            continue
        heading = re.match(r"^(#{2,4})\s+(.+)$", line)
        if heading:
            document.add_heading(heading[2], level=len(heading[1]) - 1)
        elif line.startswith("- "):
            inline(document.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"^\d+\.\s", line):
            inline(document.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s+", "", line))
        else:
            inline(document.add_paragraph(), line)
        index += 1


def main():
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.15)
    section.right_margin = Cm(2.15)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)
    normal = document.styles["Normal"]
    east_asia_font(normal)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name in ["Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
        east_asia_font(document.styles[name], "微软雅黑")
        document.styles[name].font.color.rgb = RGBColor.from_string(NAVY)
    document.styles["Heading 1"].font.size = Pt(15)
    document.styles["Heading 1"].paragraph_format.space_before = Pt(14)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(7)
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        document.styles[name].paragraph_format.keep_with_next = True
    for name in ["List Bullet", "List Number"]:
        east_asia_font(document.styles[name])
        document.styles[name].font.size = Pt(10.5)
        document.styles[name].paragraph_format.space_after = Pt(5)
    code = document.styles.add_style("Code", 1)
    east_asia_font(code)
    code.font.name = "Consolas"
    code.font.size = Pt(9)
    code.paragraph_format.line_spacing = 1.1
    code.paragraph_format.space_after = Pt(2)
    flow = document.styles.add_style("Flow", 1)
    east_asia_font(flow, "微软雅黑")
    flow.font.size = Pt(10)
    flow.font.color.rgb = RGBColor.from_string(TEAL)
    flow.paragraph_format.space_after = Pt(6)
    header = section.header.paragraphs[0]
    header.add_run("LOOMA / 灵犀胶囊").font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(GREY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("v5.0 · 2026-10-27  |  ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.title = "Looma v5 沟通收尾与项目接续全景规格书"
    document.core_properties.subject = "商业计划、产品需求与本地数据架构 v5.0"
    document.core_properties.author = "Looma Project"
    document.core_properties.comments = "由 docs 下三份 Markdown 生成；商业假设与工程验证结果分开记录。"

    document.add_paragraph("LOOMA", style="Subtitle")
    title = document.add_paragraph(style="Title")
    title.add_run("灵犀胶囊\n让客户沟通有清楚的收尾")
    document.add_paragraph("商业化、产品需求与系统技术架构", style="Subtitle")
    document.add_paragraph("版本 v5.0  /  批准 2026-10-27  /  沟通收尾与项目接续")
    document.add_paragraph("面向独立接单者、顾问与小工作室主理人。")
    document.add_paragraph("本地草稿 → 四类核对 → 确认收尾 → 下次接续 → 独立审核纪要", style="Flow")
    document.add_heading("阅读说明", level=1)
    document.add_paragraph("本文汇编三份已更新的Markdown规范，记录已批准的v5设计与商业假设。工程实现、测试和真实运行结果见implementation-v5-status.md，本规格书不表示验收已通过。")
    document.add_paragraph("首轮为本地单人真实案例试用，不提供账户支付、云同步或团队协作。生活态保留，工作态承担本轮商业验证。")
    add_table(document, [
        ["卷", "内容", "阅读目的"],
        ["第一卷", "商业化白皮书与商业计划书", "明确人群、竞争事实、价值假设与验证路径"],
        ["第二卷", "产品需求规格说明书", "明确确认流程、空间边界、提醒与外发规则"],
        ["第三卷", "系统技术架构与数据字典规范", "明确接口、迁移、凭据和工程验收边界"],
    ])
    document.add_paragraph("证据边界：网站开发者多客户场景是假设案例，待真实用户核对。竞品依据为2026-10-26公开一手资料，未穷尽全球检索。5人7天试用、目标阈值与历史试价均待验证。")
    titles = ["第一卷 / 商业化白皮书与商业计划书", "第二卷 / 产品需求规格说明书", "第三卷 / 系统技术架构与数据字典规范"]
    for source, volume in zip(SOURCES, titles):
        document.add_page_break()
        document.add_heading(volume, level=0)
        add_markdown(document, source)
    document.save(OUTPUT)
    print(f"Saved {OUTPUT}")
    print(f"Paragraphs: {len(document.paragraphs)}; tables: {len(document.tables)}")


if __name__ == "__main__":
    main()
